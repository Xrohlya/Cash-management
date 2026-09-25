from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from database.db import get_connection
from database.repository import financial_period_start, financial_period_end

BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_DIR = BASE_DIR / "data" / "reports"


def build_excel_report(user_id: int, start=None, end=None):
    start = start or financial_period_start()
    end = end or financial_period_end()
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    path = REPORT_DIR / f"budget_{user_id}_{start.isoformat()}_report.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "Операции"
    headers = ["Дата", "Тип", "Сумма", "Категория"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9EAF7")
    with get_connection() as conn:
        rows = conn.execute("SELECT created_at, kind, amount, description FROM transactions WHERE user_id=? AND created_at>=? AND created_at<? ORDER BY created_at", (user_id,start.isoformat(),end.isoformat())).fetchall()
    for r in rows:
        ws.append([r["created_at"], r["kind"], float(r["amount"]), r["description"]])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in {"A":22,"B":14,"C":14,"D":32}.items(): ws.column_dimensions[col].width = width
    cats = wb.create_sheet("Категории"); cats.append(["Категория","Сумма"])
    totals = {}
    for r in rows:
        if r["kind"] == "expense": totals[r["description"]] = totals.get(r["description"],0)+float(r["amount"])
    for k,v in sorted(totals.items(), key=lambda x:x[1], reverse=True): cats.append([k,v])
    if totals:
        chart = BarChart(); chart.title = "Расходы по категориям"; chart.y_axis.title = "Рубли"; chart.x_axis.title = "Категория"
        chart.add_data(Reference(cats,min_col=2,min_row=1,max_row=cats.max_row), titles_from_data=True)
        chart.set_categories(Reference(cats,min_col=1,min_row=2,max_row=cats.max_row)); cats.add_chart(chart,"D2")
    wb.save(path)
    return path
